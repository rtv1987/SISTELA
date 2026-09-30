import io
import json
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from sistela.api import create_app
from sistela.db import make_engine, migrate
from sistela.models import MappingConfirmation
from sistela.services import ServiceError
from sistela.spreadsheets import parse_rows, read_workbook, workbook_preview


@pytest.fixture
def client(tmp_path):
    migrate(tmp_path)
    with TestClient(create_app(tmp_path), base_url="http://127.0.0.1") as value:
        yield value


def project(client):
    return client.post('/projects', json={'name': 'Workflow', 'system_type': 'GSS'}).json()['id']


def line(client, pid, **changes):
    values = dict(project_description='Kabelio montavimas', output_description='Projekto tekstas',
                  quantity='12.345678', unit='m', system_type='GSS', line_type='Work')
    values.update(changes)
    result = client.post(f'/projects/{pid}/lines', json=values)
    assert result.status_code == 201, result.text
    return result.json()


def confirm(client, row, code='TEST-1', unit='m'):
    return client.post(f"/projects/{row['project_id']}/lines/{row['id']}/mapping/confirm", json={
        'version': row['version'], 'sistela_code': code,
        'sistela_original_description': 'Normatyvo originalas', 'normative_unit': unit})


def test_atomic_grid_conflict_delete_undo_and_exact_decimal(client):
    pid = project(client)
    a, b = line(client, pid), line(client, pid)
    endpoint = f'/projects/{pid}/grid'
    result = client.post(endpoint, json={'edits': [
        {'id': a['id'], 'version': 1, 'values': {'quantity': '0.000001'}},
        {'id': b['id'], 'version': 9, 'values': {'quantity': '2'}}]})
    assert result.status_code == 409
    assert client.get(f'/projects/{pid}/lines').json()[0]['quantity'] == '12.345678'
    result = client.post(endpoint, json={'edits': [{'id': a['id'], 'version': 1,
        'values': {'quantity': '0.000001'}}], 'deletes': [{'id': b['id'], 'version': 1}]})
    assert result.status_code == 200
    assert len(result.json()) == 1 and result.json()[0]['quantity'] == '0.000001'
    restored = client.post(f'/projects/{pid}/undo').json()
    assert len(restored) == 2 and all(r['quantity'] == '12.345678' for r in restored)
    assert client.post(endpoint, json={'edits': [{'id': a['id'], 'version': 3,
        'values': {'source_raw_text': 'tampered'}}]}).status_code == 422


def test_undo_refuses_to_overwrite_mapping_confirmation(client):
    pid = project(client)
    row = line(client, pid)
    updated = client.post(f'/projects/{pid}/grid', json={'edits': [{'id': row['id'],
        'version': 1, 'values': {'notes': 'note'}}]}).json()[0]
    assert confirm(client, updated).status_code == 200
    assert client.post(f'/projects/{pid}/undo').status_code == 409
    assert client.get(f'/projects/{pid}/lines').json()[0]['mapping_status'] == 'confirmed'


def test_duplicate_rows_and_projects_reset_approval(client):
    pid = project(client)
    row = confirm(client, line(client, pid)).json()
    copied = client.post(f'/projects/{pid}/grid', json={'duplicates': [
        {'id': row['id'], 'version': row['version']}]}).json()
    assert len(copied) == 2
    assert copied[1]['sistela_code'] == 'TEST-1' and copied[1]['mapping_status'] == 'unmapped'
    assert copied[1]['output_description'] == row['output_description']
    new = client.post(f'/projects/{pid}/duplicate').json()['id']
    clones = client.get(f'/projects/{new}/lines').json()
    assert len(clones) == 2 and all(r['mapping_status'] == 'unmapped' for r in clones)
    assert {r['id'] for r in clones}.isdisjoint(r['id'] for r in copied)


def test_history_survives_restart_and_manual_correction_wins(client, tmp_path):
    pid = project(client)
    row = confirm(client, line(client, pid)).json()
    assert confirm(client, row).json()['version'] == row['version']  # no double count
    corrected = confirm(client, row, 'TEST-2').json()
    assert corrected['confidence'] is None
    target = line(client, project(client))
    url = f"/projects/{target['project_id']}/lines/{target['id']}"
    with TestClient(create_app(tmp_path), base_url='http://127.0.0.1') as reopened:
        candidates = reopened.get(url + '/suggestions').json()
        assert candidates[0]['sistela_code'] == 'TEST-2'
        assert candidates[0]['method'] == 'exact' and candidates[0]['confirmed_count'] == 1
        applied = reopened.post(url + '/mapping/apply', json={'version': 1,
            'mapping_id': candidates[0]['mapping_id']}).json()
        assert applied['mapping_status'] == 'suggested'
        assert applied['output_description'] == 'Projekto tekstas'
    engine = make_engine(tmp_path)
    with Session(engine) as session:
        events = session.scalars(select(MappingConfirmation)).all()
        assert len(events) == 2
        assert events[1].previous_code == 'TEST-1' and events[1].selected_code == 'TEST-2'
    engine.dispose()


@pytest.mark.parametrize('description,method', [('KABELIO MONTAVIMAS', 'normalized'),
    ('Kabelio montavimo darbai', 'fuzzy')])
def test_normalized_and_fuzzy_suggestions(client, description, method):
    pid = project(client)
    confirm(client, line(client, pid))
    row = line(client, pid, project_description=description)
    candidate = client.get(f"/projects/{pid}/lines/{row['id']}/suggestions").json()[0]
    assert candidate['method'] == method and candidate['compatible']


def test_units_materials_system_isolation_and_entry_reset(client):
    pid = project(client)
    confirm(client, line(client, pid))
    row = line(client, pid, unit='100m')
    url = f"/projects/{pid}/lines/{row['id']}"
    candidates = client.get(url + '/suggestions').json()
    assert not candidates[0]['compatible'] and Decimal(candidates[0]['confidence']) < Decimal('.5')
    assert client.post(url + '/mapping/apply', json={'version': 1,
        'mapping_id': candidates[0]['mapping_id']}).status_code == 422
    assert confirm(client, row).status_code == 422
    assert client.post(url + '/entry', json={'version': 1, 'entered': True}).status_code == 422
    confirmed = confirm(client, row, unit='100m').json()
    entered = client.post(url + '/entry', json={'version': confirmed['version'], 'entered': True}).json()
    assert entered['entered_at']
    assert client.post(url + '/entry', json={'version': 1, 'entered': False}).status_code == 409
    edited = client.post(f'/projects/{pid}/grid', json={'edits': [{'id': row['id'],
        'version': entered['version'], 'values': {'output_description': 'Changed'}}]}).json()[-1]
    assert edited['entered_at'] is None and edited['mapping_status'] == 'confirmed'
    for changes in ({'line_type': 'Material'}, {'system_type': 'AS'}):
        other = line(client, pid, **changes)
        assert client.get(f"/projects/{pid}/lines/{other['id']}/suggestions").json() == []


def workbook_bytes(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Data'
    for row in rows:
        ws.append(row)
    wb.create_sheet('Empty')
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


def test_xlsx_validation_partial_import_and_selection(client):
    pid = project(client)
    data = workbook_bytes([['Name', 'Unit', 'Amount'], ['Good', '100m', '1.2345678'], ['Bad', 'm', '=1+2']])
    endpoint = f'/projects/{pid}/imports/xlsx'
    form = {'sheet': 'Data', 'header_row': 1,
            'mapping': json.dumps({'project_description': 0, 'unit': 1, 'quantity': 2})}
    assert client.post(endpoint, files={'file': ('input.xlsx', data)}, data=form).status_code == 422
    assert client.get(f'/projects/{pid}/lines').json() == []
    response = client.post(endpoint, files={'file': ('input.xlsx', data)}, data={**form, 'allow_partial': 'true'})
    assert response.status_code == 201, response.text
    run = response.json()
    assert run['rows_detected'] == 1 and {w['code'] for w in run['warnings']} == {'ROW_REJECTED', 'DECIMAL_ROUNDED'}
    row = client.get(f'/projects/{pid}/lines').json()[0]
    assert row['quantity'] == '1.234568' and row['unit'] == '100m'
    assert '1.2345678' in row['source_raw_text']
    assert client.post(endpoint, files={'file': ('input.xlsx', data)}, data={**form, 'allow_partial': 'true'}).status_code == 409
    assert client.post('/xlsx/preview', files={'file': ('input.xlsx', data)}, data={'header_row': 999}).status_code == 422


def test_xlsx_roundtrip_preserves_decimal_and_does_not_execute_text(client):
    pid = project(client)
    line(client, pid, project_description='=HYPERLINK("test")', quantity='123456789012345.123456', work_price='0.100001')
    data = client.get(f'/projects/{pid}/export.xlsx').content
    wb = load_workbook(io.BytesIO(data), data_only=False)
    assert wb.worksheets[0]['D2'].data_type == 's'
    assert wb.worksheets[0]['G2'].data_type == 's'
    wb.close()
    preview = workbook_preview(data)
    rows, warnings = parse_rows(read_workbook(data)[preview['sheet']], preview['mapping'], 1, 'GSS', preview['sheet'])
    assert not warnings and len(rows) == 1
    assert rows[0]['quantity'] == Decimal('123456789012345.123456')
    assert rows[0]['project_description'].startswith('=HYPERLINK')
    assert rows[0]['work_price'] == Decimal('0.100001')


@pytest.mark.fixtures
def test_real_xlsx_all_thirty_rows(client, source_file):
    pid = project(client)
    data = source_file('pvz.xlsx').read_bytes()
    preview = workbook_preview(data)
    response = client.post(f'/projects/{pid}/imports/xlsx', files={'file': ('pvz.xlsx', data)},
        data={'sheet': preview['sheet'], 'header_row': preview['header_row'], 'mapping': json.dumps(preview['mapping'])})
    assert response.status_code == 201, response.text
    assert response.json()['rows_detected'] == 30
    assert len(response.json()['warnings']) == 2
    rows = client.get(f'/projects/{pid}/lines').json()
    assert sum(r['line_type'] == 'Material' for r in rows) == 20
    assert sum(r['line_type'] == 'Work' for r in rows) == 10
    assert all(r['mapping_status'] == 'unmapped' for r in rows)


@pytest.mark.parametrize('data', [b'', b'not a workbook'])
def test_invalid_xlsx_is_controlled(data):
    with pytest.raises(ServiceError) as error:
        read_workbook(data)
    assert error.value.status == 422


def test_capabilities_keep_unverified_writes_blocked(client):
    items = {item['id']: item for item in client.get('/integration/capabilities').json()}
    assert items['manual_entry']['production']
    assert items['package_text_export']['status'] == 'analysis_only'
    assert items['dbf_write']['status'] == 'analysis_only'
    assert not items['dbf_write']['production']
    assert all(not item['writes_to_sistela'] for item in items.values())
