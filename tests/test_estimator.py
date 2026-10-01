"""PHASE 7 regression: source quantities, explicit review and local handoff."""

import io
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import inspect
from test_workflow import confirm, line, project

from sistela.api import create_app
from sistela.db import make_engine, migrate


@pytest.fixture
def client(tmp_path):
    migrate(tmp_path)
    with TestClient(create_app(tmp_path), base_url="http://127.0.0.1") as value:
        yield value


def review(client, row, action, **kwargs):
    return client.post(
        f"/projects/{row['project_id']}/lines/{row['id']}/review",
        json=dict(version=row["version"], action=action, **kwargs),
    )


def report(client, pid):
    response = client.post(f"/projects/{pid}/validate")
    assert response.status_code == 200, response.text
    return response.json()


def test_review_blocking_warning_and_empty_project(client):
    pid = project(client)
    assert report(client, pid)["blocking"] == 1
    material = line(client, pid, line_type="Material", quantity="3", unit="vnt.")
    result = report(client, pid)
    assert result["blocking"] == 0 and result["warnings"] == 1
    assert result["status"] == "READY_FOR_SISTELA"
    assert result["rows"][0]["issues"][0]["category"] == "PRICE_MISSING"
    assert review(client, material, "convert", target_unit="100m").status_code == 422
    row = line(client, pid, quantity="0", unit="unknown")
    result = report(client, pid)
    categories = {
        i["category"] for r in result["rows"] if r["id"] == row["id"] for i in r["issues"]
    }
    assert {
        "MISSING_SISTELA_CODE",
        "MISSING_QUANTITY",
        "UNKNOWN_UNIT",
    } <= categories


@pytest.mark.parametrize("path", ["grid", "put"])
def test_explicit_conversion_stales_on_source_change(client, path):
    pid = project(client)
    row = line(client, pid, quantity="650")
    assert confirm(client, row, unit="100M").status_code == 422
    row = review(client, row, "convert", target_unit="100M").json()
    conversion = row["review_data"]["conversion"]
    assert conversion["confirmed_by_user"] is True
    assert Decimal(conversion["target_quantity"]) == Decimal("6.5")
    assert row["quantity"] == "650" and row["unit"] == "m"
    row = confirm(client, row, unit="100M").json()
    assert report(client, pid)["status"] == "READY_FOR_SISTELA"
    assert client.get(f"/projects/{pid}").json()["status"] == "READY_FOR_SISTELA"
    if path == "grid":
        response = client.post(
            f"/projects/{pid}/grid",
            json={
                "edits": [
                    {"id": row["id"], "version": row["version"], "values": {"quantity": "700"}}
                ]
            },
        )
    else:
        from sistela.schemas import LineCreate

        response = client.put(
            f"/projects/{pid}/lines/{row['id']}",
            json={
                **{k: row[k] for k in LineCreate.model_fields},
                "quantity": "700",
                "version": row["version"],
            },
        )
    assert response.status_code == 200, response.text
    current = client.get(f"/projects/{pid}/lines").json()[0]
    assert current["mapping_status"] == "needs_review"
    assert current["review_data"]["conversion"]["confirmed_by_user"] is False
    assert report(client, pid)["blocking"] > 0
    assert confirm(client, current, unit="100m").status_code == 422
    assert (
        client.post(
            f"/projects/{pid}/lines/{row['id']}/entry",
            json={"version": current["version"], "entered": True},
        ).status_code
        == 422
    )


def test_conversion_rules_versions_and_reconfirmation(client):
    row = line(client, project(client), quantity="0.000001", unit="100M")
    assert review(client, row, "convert", target_unit="kg").status_code == 422
    converted = review(client, row, "convert", target_unit="m").json()
    assert Decimal(converted["review_data"]["conversion"]["target_quantity"]) == Decimal(".0001")
    assert review(client, row, "convert", target_unit="m").status_code == 409
    assert Decimal(converted["quantity"]) == Decimal(".000001")
    confirmed = confirm(client, converted, unit="m").json()
    restored = review(client, confirmed, "cancel_conversion").json()
    assert "conversion" not in restored["review_data"]
    assert restored["unit"] == "100M" and restored["quantity"] == converted["quantity"]
    assert restored["mapping_status"] == "needs_review"
    assert confirm(client, restored, unit="100M").status_code == 200


def test_lifecycle_rejection_manual_and_metrics(client):
    pid = project(client)
    seed = confirm(client, line(client, pid)).json()
    row = line(client, pid)
    url = f"/projects/{pid}/lines/{row['id']}"
    suggestion = client.get(url + "/suggestions").json()[0]
    row = client.post(
        url + "/mapping/apply",
        json={"version": row["version"], "mapping_id": suggestion["mapping_id"]},
    ).json()
    assert row["mapping_status"] == "suggested"
    row = confirm(client, row).json()
    assert report(client, pid)["statistics"]["confirmed_without_change_count"] == 1
    row = review(client, row, "reject").json()
    assert row["mapping_status"] == "rejected" and row["sistela_code"] == ""
    assert suggestion["mapping_id"] not in [
        s["mapping_id"] for s in client.get(url + "/suggestions").json()
    ]
    row = review(client, row, "manual").json()
    assert row["review_data"]["manual"] and row["mapping_status"] == "needs_review"
    assert report(client, pid)["blocking"] > 0
    row = confirm(client, row, code="MANUAL-2").json()
    assert row["mapping_status"] == "confirmed"
    assert report(client, pid)["statistics"]["manual_mapping_count"] == 2
    assert seed["source_raw_text"] == row["source_raw_text"]


def test_duplicate_warning_acknowledgement_and_price_status(client):
    pid = project(client)
    a = line(client, pid, line_type="Material", material_price="2.10")
    b = line(client, pid, line_type="Material", material_price="2.10")
    assert any(
        i["category"] == "POSSIBLE_DUPLICATE"
        for r in report(client, pid)["rows"]
        for i in r["issues"]
    )
    a = review(client, a, "keep_both", duplicate_id=b["id"]).json()
    assert not any(
        i["category"] == "POSSIBLE_DUPLICATE"
        for r in report(client, pid)["rows"]
        for i in r["issues"]
    )
    a = review(client, a, "price", price_status="CONFIRMED").json()
    assert report(client, pid)["summary"]["confirmed_prices"] == 1
    assert a["material_price"] == "2.10"
    client.post(
        f"/projects/{pid}/grid",
        json={
            "edits": [{"id": a["id"], "version": a["version"], "values": {"material_price": "3"}}]
        },
    )
    assert report(client, pid)["summary"]["confirmed_prices"] == 0
    current = client.get(f"/projects/{pid}/lines").json()[0]
    client.post(
        f"/projects/{pid}/grid",
        json={
            "edits": [
                {
                    "id": current["id"],
                    "version": current["version"],
                    "values": {"material_price": None},
                }
            ]
        },
    )
    assert report(client, pid)["summary"]["missing_prices"] == 1


def test_working_export_and_reload(client):
    pid = project(client)
    row = review(client, line(client, pid, quantity="650"), "convert", target_unit="100m").json()
    row = confirm(client, row, unit="100m").json()
    response = client.post(
        f"/projects/{pid}/lines/{row['id']}/entry",
        json={"version": row["version"], "entered": True},
    )
    assert response.status_code == 200
    assert client.get(f"/projects/{pid}/lines").json()[0]["entered_at"]
    result = report(client, pid)
    assert result["statistics"]["entry_mode_completed_count"] == 1
    assert result["status"] == "READY_FOR_SISTELA"  # Manual progress never defines product completion.
    wb = load_workbook(io.BytesIO(client.get(f"/projects/{pid}/export.xlsx").content))
    rows = list(wb["Darbo lentelė"].values)
    exported = dict(zip(rows[0], rows[1]))
    assert rows[1][6] == 650
    assert exported["SISTELA tikslinis kiekis"] == 6.5
    assert exported["SISTELA tikslinis vienetas"] == "100m"
    assert exported["Normatyvo būsena"] == "confirmed"
    wb.close()


def test_migration_preserves_checks(tmp_path):
    migrate(tmp_path)
    engine = make_engine(tmp_path)
    constraints = inspect(engine).get_check_constraints("estimate_lines")
    assert len(constraints) == 3
    assert any("needs_review" in c["sqltext"] for c in constraints)
    assert inspect(engine).get_foreign_keys("estimate_lines")
    engine.dispose()


def test_conversion_unlocks_suggestion_without_confirming_or_changing_source(client):
    pid = project(client)
    confirm(client, line(client, pid, unit="100m", quantity="6.5"), unit="100m")
    row = line(client, pid, quantity="650")
    url = f"/projects/{pid}/lines/{row['id']}"
    suggestion = client.get(url + "/suggestions").json()[0]
    assert not suggestion["compatible"]
    assert (
        client.post(
            url + "/mapping/apply",
            json={"version": row["version"], "mapping_id": suggestion["mapping_id"]},
        ).status_code
        == 422
    )
    row = review(client, row, "convert", target_unit="100M").json()
    row = client.post(
        url + "/mapping/apply",
        json={"version": row["version"], "mapping_id": suggestion["mapping_id"]},
    ).json()
    assert row["quantity"] == "650" and row["unit"] == "m" and row["mapping_status"] == "suggested"
    assert report(client, pid)["blocking"] == 0  # Phase 10: selected code needs no second confirmation.
    row = confirm(client, row, code="CHANGED", unit="100m").json()
    assert report(client, pid)["statistics"]["changed_mapping_count"] == 1
    assert report(client, pid)["blocking"] == 0
    duplicate = client.post(
        f"/projects/{pid}/grid", json={"duplicates": [{"id": row["id"], "version": row["version"]}]}
    ).json()[-1]
    assert duplicate["review_data"] == {} and duplicate["mapping_status"] == "unmapped"


def test_reject_incompatible_suggestion_without_applying(client):
    pid = project(client)
    confirm(client, line(client, pid, unit="100m"), unit="100m")
    row = line(client, pid)
    url = f"/projects/{pid}/lines/{row['id']}"
    suggestion = client.get(url + "/suggestions").json()[0]
    row = review(client, row, "reject", mapping_id=suggestion["mapping_id"]).json()
    assert row["mapping_status"] == "rejected"
    assert not client.get(url + "/suggestions").json()
    assert row["unit"] == "m" and row["quantity"] == "12.345678"


def test_material_price_reference_is_only_a_warning(client):
    pid = project(client)
    row = line(client, pid, line_type="Material")
    row = review(client, row, "price", price_status="HISTORICAL_REFERENCE").json()
    assert row["material_price"] is None
    result = report(client, pid)
    assert result["blocking"] == 0
    assert result["summary"]["missing_prices"] == 1
    assert {"PRICE_MISSING", "HISTORICAL_PRICE"} == {
        i["category"] for i in result["rows"][0]["issues"]
    }
    assert review(client, row, "price", price_status="CONFIRMED").status_code == 422


def test_duplicate_requires_same_system_unit_and_type(client):
    pid = project(client)
    line(client, pid, line_type="Material", unit="m")
    line(client, pid, line_type="Material", unit="m", system_type="AS")
    line(client, pid, line_type="Material", unit="vnt.")
    line(client, pid, line_type="Work", unit="m")
    assert not any(
        i["category"] == "POSSIBLE_DUPLICATE"
        for r in report(client, pid)["rows"]
        for i in r["issues"]
    )


def test_conversion_uses_target_unit_for_every_candidate(client):
    pid = project(client)
    confirm(client, line(client, pid), code="METRES", unit="m")
    confirm(client, line(client, pid, unit="100m"), code="HUNDREDS", unit="100m")
    row = review(client, line(client, pid, quantity="650"), "convert", target_unit="100m").json()
    url = f"/projects/{pid}/lines/{row['id']}"
    candidates = client.get(url + "/suggestions").json()
    assert candidates[0]["sistela_code"] == "HUNDREDS" and candidates[0]["compatible"]
    metre = next(s for s in candidates if s["sistela_code"] == "METRES")
    assert not metre["compatible"]
    assert (
        client.post(
            url + "/mapping/apply",
            json={"version": row["version"], "mapping_id": metre["mapping_id"]},
        ).status_code
        == 422
    )


def test_manual_confirmation_does_not_count_its_own_new_suggestion(client):
    pid = project(client)
    row = confirm(client, line(client, pid, project_description="Previously unseen task")).json()
    assert report(client, pid)["statistics"]["auto_suggested_mapping_count"] == 0
    assert report(client, pid)["statistics"]["manual_mapping_count"] == 1
    assert row["review_data"]["suggestion_offered"] is False


def test_tiny_target_export_is_not_displayed_as_zero(client):
    pid = project(client)
    row = review(
        client, line(client, pid, quantity="0.000001"), "convert", target_unit="100m"
    ).json()
    assert Decimal(row["review_data"]["conversion"]["target_quantity"]) == Decimal(".00000001")
    wb = load_workbook(io.BytesIO(client.get(f"/projects/{pid}/export.xlsx").content))
    cell = wb["Darbo lentelė"]["P2"]
    assert Decimal(str(cell.value)) == Decimal(".00000001")
    assert cell.number_format == "0.########"
    wb.close()


def test_real_gss_review_acceptance_and_provenance(client, source_file):
    files = {
        role + "25-04-14.dbf": source_file(role + "25-04-14.dbf").read_bytes()
        for role in ("sd", "dd", "nd", "pd", "td", "od")
    }
    assert (
        client.post(
            "/history/imports",
            files=[("files", (name, data)) for name, data in files.items()],
            data={"encoding": "cp1257"},
        ).status_code
        == 200
    )
    pid = project(client)
    pdf = source_file("2024-10-XX-TDP-GSS.pdf")
    result = client.post(
        f"/projects/{pid}/imports/pdf",
        files={"file": (pdf.name, pdf.read_bytes(), "application/pdf")},
    )
    assert result.json()["rows_detected"] == 21
    result = report(client, pid)
    assert result["summary"]["materials"] == 11 and result["summary"]["works"] == 10
    suggested = sum(bool(r["suggestions"]) for r in result["rows"])
    compatible = sum(any(s["compatible"] for s in r["suggestions"]) for r in result["rows"])
    print(
        f"GSS_MEASUREMENT suggestions={suggested}/10 compatible={compatible}/10 manual_review=10/10"
    )
    assert suggested > 0
    originals = client.get(f"/projects/{pid}/lines").json()
    for row in originals:
        if row["line_type"] != "Work":
            continue
        unit = row["unit"]
        if unit == "m":
            row = review(client, row, "convert", target_unit="100m").json()
            unit = "100m"
        # Explicit synthetic operator decisions, never production hardcoded codes.
        assert confirm(client, row, code="TEST-HANDOFF", unit=unit).status_code == 200
    result = report(client, pid)
    assert result["status"] == "READY_FOR_SISTELA" and result["blocking"] == 0
    current = client.get(f"/projects/{pid}/lines").json()
    assert [
        (r["quantity"], r["unit"], r["source_page"], r["source_raw_text"], r["source_document_id"])
        for r in current
    ] == [
        (r["quantity"], r["unit"], r["source_page"], r["source_raw_text"], r["source_document_id"])
        for r in originals
    ]
    assert all(r["source_filename"] == pdf.name for r in result["rows"])
