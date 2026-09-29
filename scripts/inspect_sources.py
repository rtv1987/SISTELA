"""Read-only source inventory. Writes only to tmp/source-inspection/."""

import hashlib
import json
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "tmp" / "source-inspection"


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    inventory = []
    for path in sorted(ROOT.iterdir()):
        if path.suffix.lower() not in {".pdf", ".xlsx", ".dbf"}:
            continue
        data = path.read_bytes()
        entry = {
            "filename": path.name,
            "size": len(data),
            "sha256": hashlib.sha256(data).hexdigest(),
        }
        if path.suffix.lower() == ".pdf":
            with pdfplumber.open(path) as pdf:
                entry["pages"] = len(pdf.pages)
                texts = []
                for number, page in enumerate(pdf.pages, 1):
                    text = page.extract_text() or ""
                    texts.append(f"\n--- PAGE {number} ---\n{text}")
                (OUT / (path.stem + ".txt")).write_text("\n".join(texts), encoding="utf-8")
        elif path.suffix.lower() == ".xlsx":
            wb = load_workbook(path, read_only=False, data_only=False)
            entry["sheets"] = [
                {
                    "name": s.title,
                    "rows": s.max_row,
                    "columns": s.max_column,
                    "merged_ranges": [str(r) for r in s.merged_cells.ranges],
                    "formula_count": sum(c.data_type == "f" for row in s for c in row),
                }
                for s in wb
            ]
            rows = {
                s.title: [
                    {"row": i, "cells": {str(j): v for j, v in enumerate(row, 1) if v is not None}}
                    for i, row in enumerate(s.iter_rows(values_only=True), 1)
                    if any(v is not None for v in row)
                ]
                for s in wb
            }
            (OUT / "xlsx-rows.json").write_text(
                json.dumps(rows, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
            )
            wb.close()
        else:
            entry.update(
                version=data[0],
                records=int.from_bytes(data[4:8], "little"),
                header_length=int.from_bytes(data[8:10], "little"),
                record_length=int.from_bytes(data[10:12], "little"),
                language_driver=data[29],
            )
            fields = []
            offset = 32
            while offset < entry["header_length"] and data[offset] != 13:
                field = data[offset : offset + 32]
                fields.append(
                    {
                        "name": field[:11].split(b"\0")[0].decode("ascii"),
                        "type": chr(field[11]),
                        "length": field[16],
                        "decimals": field[17],
                    }
                )
                offset += 32
            entry["fields"] = fields
        inventory.append(entry)
    (OUT / "inventory.json").write_text(
        json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(
        json.dumps(
            [
                {k: v for k, v in item.items() if k not in {"fields", "sheets"}}
                for item in inventory
            ],
            ensure_ascii=True,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
