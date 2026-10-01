"""Local XLSX exchange. No claim of SISTELA compatibility, no formula execution."""
import hashlib
import io
import json
import posixpath
import time
import zipfile
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path, PureWindowsPath
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select

from .models import EstimateLine, EstimateSection, ImportRun, SourceDocument, utc_now
from .parsers.common import compact, normalize_unit, parse_decimal
from .ports import ExportResult
from .schemas import LineCreate
from .services import ServiceError, get_project, logger, next_order, store_document

FIELDS = {"source_position": "Pozicija", "project_description": "Projekto pavadinimas",
    "technical_reference": "Modelis / žymuo", "unit": "Vnt.", "quantity": "Kiekis",
    "sistela_code": "SISTELA kodas", "sistela_original_description": "Originalus normatyvo pavadinimas",
    "output_description": "Galutinis pavadinimas", "material_price": "Medžiagos kaina",
    "work_price": "Darbo kaina", "unit_price": "Vieneto kaina (pagal tipą)", "notes": "Pastabos",
    "line_type": "Tipas", "system_type": "Sistema"}
ALIASES = {
    "source_position": {"pozicija", "eilnr", "sameil", "nr"},
    "project_description": {"projektopavadinimas", "pavadinimas", "darbuirislaiduaprasymai", "aprasymas"},
    "technical_reference": {"modelis", "modeliszymuo", "zymuo"}, "unit": {"vnt", "matovnt", "vienetas"},
    "quantity": {"kiekis", "quantity"}, "sistela_code": {"sistelakodas", "darbokodas", "kodas"},
    "sistela_original_description": {"originalusnormatyvopavadinimas"}, "output_description": {"galutinispavadinimas"},
    "material_price": {"medziagoskaina", "medziagoskainaeur"}, "work_price": {"darbokaina", "darbokainaeur"},
    "unit_price": {"vienetokaina", "vienetokainaeur"}, "notes": {"pastabos"},
    "line_type": {"tipas", "type"}, "system_type": {"sistema", "system"},
}


def read_workbook(data: bytes) -> dict[str, list[list[str | None]]]:
    if len(data) > 25 * 1024 * 1024:
        raise ServiceError(413, "XLSX limitas 25 MB.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if len(archive.infolist()) > 5000 or sum(i.file_size for i in archive.infolist()) > 80 * 1024 * 1024:
                raise ServiceError(422, "XLSX išpakuotas turinys per didelis.")
            ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
            rels = {r.attrib["Id"]: r.attrib["Target"] for r in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))}
            sheets = ET.fromstring(archive.read("xl/workbook.xml")).find("s:sheets", ns)
            numeric = {}
            for sheet in sheets:
                target = rels[sheet.attrib[f"{{{rel_ns}}}id"]]
                target = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
                cells = ET.fromstring(archive.read(target)).findall(".//s:c", ns)
                if len(cells) > 100000:
                    raise ServiceError(422, "Per daug langelių. Naudokite mažesnį darbo lapą.")
                numeric[sheet.attrib["name"]] = {c.attrib["r"]: c.find("s:v", ns).text
                    for c in cells if c.attrib.get("t", "n") == "n" and c.find("s:v", ns) is not None}
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            result = {}
            for sheet in workbook:
                if (sheet.max_row or 0) > 10000 or (sheet.max_column or 0) > 80:
                    raise ServiceError(422, "Lapas viršija 10000 eilučių arba 80 kolonų ribą.")
                result[sheet.title] = [[None if cell.value is None else
                    "#FORMULA_NOT_EVALUATED" if cell.data_type == "f" else
                    numeric[sheet.title].get(cell.coordinate, str(cell.value)) for cell in row]
                    for row in sheet.iter_rows()]
            return result
        finally:
            workbook.close()
    except ServiceError:
        raise
    except Exception as exc:
        raise ServiceError(422, "XLSX nepavyko perskaityti. Pasirinkite galiojantį .xlsx failą.") from exc


def guess_columns(rows, header_row=None):
    best = ({}, 1, 0)
    candidates = [header_row - 1] if header_row else range(min(len(rows), 50))
    for i in candidates:
        if i < 0 or i >= len(rows):
            continue
        for height in (1, 2):
            mapping = {}
            for j in range(len(rows[i])):
                name = compact(" ".join(str(rows[k][j] or "") for k in range(i, min(i+height, len(rows))) if j < len(rows[k])))
                for field, aliases in ALIASES.items():
                    if name in aliases:
                        mapping[field] = j
            score = len(mapping) + 5 * len({"project_description", "unit", "quantity"} & mapping.keys())
            if score > best[2]:
                best = (mapping, i + 1, score)
    return best[0], best[1]


def workbook_preview(data, sheet_name=None, header_row=None):
    sheets = read_workbook(data)
    selected = sheet_name or next(iter(sheets))
    if selected not in sheets:
        raise ServiceError(422, "Lapas nerastas.")
    rows = sheets[selected]
    if header_row is not None and not 1 <= header_row <= len(rows):
        raise ServiceError(422, "Neteisinga antraštės eilutė.")
    mapping, header = guess_columns(rows, header_row)
    return {"sheets": [{"name": name, "rows": len(values)} for name, values in sheets.items()],
            "sheet": selected, "header_row": header, "mapping": mapping, "fields": FIELDS,
            "columns": [{"index": j, "label": f"{j+1}: {cell or '—'}"} for j, cell in enumerate(rows[header-1] if rows else [])],
            "preview": rows[header:header+10], "row_count": len(rows), "checksum": hashlib.sha256(data).hexdigest()}


def parse_rows(rows, mapping, header_row, system, sheet_name):
    if not {"project_description", "unit", "quantity"} <= mapping.keys():
        raise ServiceError(422, "Susiekite pavadinimo, vieneto ir kiekio kolonas.")
    if not set(mapping) <= FIELDS.keys() or any(not isinstance(v, int) or isinstance(v, bool) or v < 0 or v >= 80 for v in mapping.values()):
        raise ServiceError(422, "Neteisingas kolonų susiejimas.")
    if len(set(mapping.values())) != len(mapping):
        raise ServiceError(422, "Viena kolona negali būti susieta su keliais laukais.")
    if header_row < 0 or header_row > len(rows):
        raise ServiceError(422, "Neteisinga antraštės eilutė.")
    result, warnings = [], []
    category = "Other"
    kinds = {"material": "Material", "medziaga": "Material", "medziagos": "Material", "work": "Work", "darbas": "Work", "darbai": "Work", "montavimodarbai": "Work", "equipment":"Equipment", "irenginys":"Equipment", "irenginiai":"Equipment", "other": "Other", "kita": "Other"}
    for index, row in enumerate(rows[header_row:], header_row+1):
        def cell(field):
            at = mapping.get(field)
            return str(row[at] or "").strip() if at is not None and at < len(row) else ""
        description = cell("project_description")
        label = compact(description)
        if label in kinds and not cell("quantity"):
            category = kinds[label]
            continue
        if not description or label in {"aprasymai", "pavadinimas", "projektopavadinimas"} or (not cell("quantity") and not cell("unit")):
            continue
        try:
            precision_warnings = []

            def number(field):
                raw = cell(field)
                try:
                    value = parse_decimal(raw)
                except ValueError:
                    # OOXML numbers can use scientific notation. Never pass through float.
                    if "e" not in raw.lower():
                        raise
                    value = Decimal(raw)
                if not value.is_finite() or value < 0 or value >= Decimal("1e18"):
                    raise ValueError("Number outside supported range")
                rounded = value.quantize(Decimal("0.000001"), rounding=ROUND_HALF_UP)
                if rounded != value:
                    precision_warnings.append({"code": "DECIMAL_ROUNDED", "row": index,
                        "message": f"{index} eilutė, {FIELDS[field]}: {raw} → {rounded}. Tikslumas: 6 skaitmenys po kablelio; originalas išsaugotas."})
                return rounded

            raw_type = compact(cell("line_type"))
            kind = kinds[raw_type] if raw_type else category
            values = {field: cell(field) for field in mapping if field in LineCreate.model_fields}
            values.update(project_description=description, output_description=cell("output_description") or description,
                          quantity=number("quantity"), unit=normalize_unit(cell("unit")),
                          system_type=cell("system_type") or system, line_type=kind)
            if not values["unit"]:
                raise ValueError("Missing unit")
            for price in ("material_price", "work_price"):
                values[price] = number(price) if cell(price) else None
            if cell("unit_price"):
                if kind == "Other":
                    raise ValueError("Unknown type for unit price")
                values["material_price" if kind in {"Material","Equipment"} else "work_price"] = number("unit_price")
            if any("#FORMULA_NOT_EVALUATED" in str(v) for v in values.values()):
                raise ValueError("Formula not evaluated")
            validated = LineCreate.model_validate(values)
            result.append({**validated.model_dump(), "source_position": cell("source_position") or str(index),
                "source_raw_text": json.dumps({"sheet": sheet_name, "row": index, "cells": row}, ensure_ascii=False)})
            warnings.extend(precision_warnings)
        except (ValueError, KeyError, InvalidOperation):
            warnings.append({"code": "ROW_REJECTED", "row": index,
                             "message": f"{index} eilutė: patikrinkite skaičius, vienetą ir tipą. Formulės nevykdomos."})
    if len(result) > 500:
        raise ServiceError(422, "Vieno importo limitas 500 eilučių.")
    if not result:
        raise ServiceError(422, "Nerasta tinkamų eilučių. Patikrinkite lapą ir kolonų susiejimą.")
    return result, warnings


def import_xlsx(session, directory, project_id, filename, data, sheet, header_row, mapping, allow_partial=False):
    project = get_project(session, project_id)
    sheets = read_workbook(data)
    if sheet not in sheets:
        raise ServiceError(422, "Lapas nerastas.")
    parsed, warnings = parse_rows(sheets[sheet], mapping, header_row, project.system_type, sheet)
    if any(w["code"] == "ROW_REJECTED" for w in warnings) and not allow_partial:
        raise ServiceError(422, "Yra netinkamų eilučių: " + " ".join(w["message"] for w in warnings[:5]) + " Dalinį importą pasirinkite aiškiai.")
    checksum = hashlib.sha256(data).hexdigest()
    selection = hashlib.sha256(json.dumps([sheet, header_row, mapping], sort_keys=True).encode()).hexdigest()
    document = session.scalar(select(SourceDocument).where(SourceDocument.project_id == project_id, SourceDocument.checksum == checksum))
    if document:
        runs = session.scalars(select(ImportRun).where(ImportRun.source_document_id == document.id)).all()
        if any(r.options.get("selection_key") == selection and r.status != "failed" for r in runs):
            raise ServiceError(409, "Šis lapas su tokiu kolonų susiejimu jau importuotas.")
    else:
        destination = directory / "documents" / f"{checksum}.xlsx"
        store_document(destination, data, checksum)
        document = SourceDocument(project_id=project_id, filename=PureWindowsPath(filename).name,
            file_type="xlsx", checksum=checksum, storage_path=destination.name)
        session.add(document)
        session.flush()
    run = ImportRun(project_id=project_id, source_document_id=document.id, parser_version="xlsx-table-1",
        options={"selection_key": selection, "sheet": sheet, "header_row": header_row, "mapping": mapping})
    session.add(run)
    session.commit()
    run_id = run.id
    started = time.perf_counter()
    logger.info("import_start id=%s parser=xlsx-table-1", run_id)
    try:
        order = next_order(session, project_id)
        sections = {}
        for i, item in enumerate(parsed):
            kind = item["line_type"]
            if kind not in sections:
                section = EstimateSection(project_id=project_id, name=kind, sort_order=order+i)
                session.add(section)
                session.flush()
                sections[kind] = section.id
            session.add(EstimateLine(project_id=project_id, section_id=sections[kind],
                source_document_id=document.id, sort_order=order+i, **item))
        run.status = "needs_review"
        run.rows_detected = len(parsed)
        run.warnings = warnings
        run.completed_at = utc_now()
        run.elapsed_ms = int((time.perf_counter()-started)*1000)
        session.flush()
        from .auto_mapping import populate_automatic
        populate_automatic(session, project_id)
        project.status = "IMPORTED"
        project.updated_at = utc_now()
        session.commit()
    except Exception:
        session.rollback()
        run = session.get(ImportRun, run_id)
        run.status = "failed"
        run.error_message = "IMPORT_FAILED: Excel importo išsaugoti nepavyko."
        run.completed_at = utc_now()
        session.commit()
    session.refresh(run)
    logger.info("import_end id=%s rows=%s status=%s", run.id, run.rows_detected, run.status)
    return run


class ExcelExporter:
    def render(self, project, lines, sections=None):
        from .handoff import target
        wb = Workbook()
        ws = wb.active
        ws.title = "Darbo lentelė"
        fields = ["system_type", "line_type", "project_description", "technical_reference", "unit", "quantity",
                  "sistela_code", "sistela_original_description", "output_description", "material_price", "work_price", "notes"]
        ws.append(["Projektas"] + [FIELDS[f] for f in fields] + ["Skyrius", "SISTELA tikslinis vienetas", "SISTELA tikslinis kiekis", "Normatyvo būsena", "Kainos būsena", "Konversija patvirtinta"])
        for line in lines:
            quantity, unit, valid = target(line)
            review = line.review_data or {}
            ws.append([project.name] + [getattr(line, f) for f in fields] + [
                (sections or {}).get(line.section_id, ""), unit if valid else "", quantity if valid else None,
                line.mapping_status, review.get("price_status", "ENTERED" if (line.material_price if line.line_type in {"Material","Equipment"} else line.work_price) is not None else "MISSING"),
                "TAIP" if review.get("conversion") and valid else "NE" if review.get("conversion") else "NEREIKALINGA"])
        warnings = []
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Never execute user descriptions as formulas.
                elif isinstance(cell.value, Decimal):
                    if len(cell.value.normalize().as_tuple().digits) > 15:
                        cell.value = format(cell.value, "f")
                        cell.data_type = "s"
                        warnings.append("High precision decimal preserved as text")
                    else:
                        cell.number_format = "0.########"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="276F4C")
        from openpyxl.utils import get_column_letter
        for i, width in enumerate([28, 12, 12, 48, 23, 12, 16, 18, 48, 48, 18, 18, 40, 24, 20, 22, 20, 22, 24], 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        ws.row_dimensions[1].height = 32
        info = wb.create_sheet("Apie eksportą")
        info.append(["SISTELA Assistant darbo eksportas"])
        info.append(["Tai nėra patvirtintas SISTELA importo formatas."])
        info.append(["Darbinė sąmata, ne galutinė komercinė SISTELA sąmata. Projekto ir patvirtinti tiksliniai kiekiai atskiri."])
        info.append(["Normatyvų ir suvedimo patvirtinimai neperkeliami per XLSX."])
        for warning in sorted(set(warnings)):
            info.append([warning])
        info.column_dimensions["A"].width = 95
        output = io.BytesIO()
        wb.save(output)
        wb.close()
        return output.getvalue(), tuple(sorted(set(warnings)))

    def export(self, project, lines, destination: Path):
        if destination.suffix.lower() != ".xlsx":
            raise ValueError("Excel destination must end in .xlsx")
        data, warnings = self.render(project, lines)
        with destination.open("xb") as output:
            output.write(data)
        return ExportResult(destination, warnings)
